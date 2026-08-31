"""材料提取服务测试（T049，D024）：类型识别 / 分派路径 / 降级兜底。

不依赖真实 LLM 与真实 PDF 渲染：提取/解析函数均 monkeypatch；
docx 用 python-docx 现场构造真实文件验证文本抽取。
"""

from __future__ import annotations

import io

from services.materials import _fields_from, detect_material_type, extract_material

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _make_docx() -> bytes:
    """构造含诊断信息的真实 .docx（正文 + 表格）。"""
    import docx

    document = docx.Document()
    document.add_paragraph("诊断证明书")
    document.add_paragraph("患者姓名：张三    临床诊断：急性阑尾炎（K35）")
    document.add_paragraph("住院费用总额：15,800.00 元    出院日期：2026-08-20")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "科室"
    table.rows[0].cells[1].text = "普外科"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ---------- 类型识别 ----------


def test_detect_by_mime() -> None:
    assert detect_material_type("a.png", "image/png") == "image"
    assert detect_material_type("a.pdf", "application/pdf") == "pdf"
    assert detect_material_type("a.docx", _DOCX_MIME) == "docx"


def test_detect_by_extension_fallback() -> None:
    """MIME 缺失/generic 时按扩展名兜底（大小写不敏感）。"""
    assert detect_material_type("证明.PDF", "") == "pdf"
    assert detect_material_type("证明.docx", "application/octet-stream") == "docx"
    assert detect_material_type("照片.JPG", "") == "image"


def test_detect_rejects_doc_and_unknown() -> None:
    """.doc 旧格式与未知类型不识别（API 层给 422 与转存提示）。"""
    assert detect_material_type("a.doc", "") == ""
    assert detect_material_type("a.doc", "application/msword") == ""
    assert detect_material_type("a.exe", "application/octet-stream") == ""


def test_amount_normalization() -> None:
    """金额宽松归一化（复用 OCR 工具的规则）。"""
    assert _fields_from({"amount": "15,800.00"})["amount"] == 15800.0
    assert _fields_from({"amount": "15800 元"})["amount"] == 15800.0
    assert _fields_from({"amount": "无"})["amount"] is None


# ---------- 提取分派 ----------


async def test_image_dispatch_uses_ocr_tool(monkeypatch) -> None:
    """图片走既有 OcrExtractTool 链路（vision + 内置 Mock 兜底）。"""
    from services import materials

    async def fake_tool(content: bytes, mime: str) -> dict:
        assert mime == "image/jpeg"
        return {
            "patient_name": "李四",
            "diagnosis": None,
            "amount": None,
            "date": None,
            "source": "vision",
        }

    monkeypatch.setattr(materials, "_extract_image_via_tool", fake_tool)
    result = await extract_material("x.jpg", "image/jpeg", b"jpeg-bytes")
    assert result.file_type == "image"
    assert result.source == "vision"
    assert result.patient_name == "李四"


async def test_docx_text_model_path(monkeypatch) -> None:
    """docx：抽正文+表格文本 → 主链路模型提取。"""
    from services import materials

    captured: dict[str, str] = {}

    async def fake_text_model(text: str) -> dict:
        captured["text"] = text
        return {
            "patient_name": "张三",
            "diagnosis": "急性阑尾炎",
            "amount": 15800.0,
            "date": "2026-08-20",
        }

    monkeypatch.setattr(materials, "_extract_via_text_model", fake_text_model)
    result = await extract_material("诊断证明.docx", _DOCX_MIME, _make_docx())
    assert result.file_type == "docx"
    assert result.source == "text_model"
    assert result.diagnosis == "急性阑尾炎"
    assert result.amount == 15800.0
    assert "张三" in captured["text"] and "普外科" in captured["text"]  # 正文与表格都抽到


async def test_docx_extraction_failure_falls_back_to_mock(monkeypatch) -> None:
    """模型提取失败 → Mock 兜底，不抛错。"""
    from services import materials

    async def fail(text: str) -> None:
        return None

    monkeypatch.setattr(materials, "_extract_via_text_model", fail)
    result = await extract_material("x.docx", _DOCX_MIME, _make_docx())
    assert result.source == "mock_fallback"
    assert result.file_type == "docx"


async def test_pdf_with_text_uses_text_model(monkeypatch) -> None:
    """文本型 PDF：文本充足走主链路模型，不触发渲染。"""
    from services import materials

    monkeypatch.setattr(
        materials, "_pdf_text", lambda content: "诊断证明 急性阑尾炎 住院费用 15800 元 患者张三" * 2
    )
    render_calls: list[int] = []

    def fake_render(content: bytes, max_pages: int) -> list[bytes]:
        render_calls.append(1)
        return []

    monkeypatch.setattr(materials, "_pdf_page_images", fake_render)

    async def fake_text_model(text: str) -> dict:
        return {"diagnosis": "急性阑尾炎", "amount": 15800.0}

    monkeypatch.setattr(materials, "_extract_via_text_model", fake_text_model)
    result = await extract_material("a.pdf", "application/pdf", b"%PDF-fake")
    assert result.source == "text_model"
    assert result.file_type == "pdf"
    assert not render_calls


async def test_pdf_scanned_renders_pages_until_hit(monkeypatch) -> None:
    """扫描件：文本不足 → 逐页 vision，取首个有效页（第 2 页命中）。"""
    from services import materials

    monkeypatch.setattr(materials, "_pdf_text", lambda content: "   ")  # 无有效文本
    monkeypatch.setattr(materials, "_pdf_page_images", lambda c, m: [b"page1", b"page2"])

    async def fake_vision(png: bytes) -> dict | None:
        return {"diagnosis": "骨折"} if png == b"page2" else None

    monkeypatch.setattr(materials, "_vision_extract_png", fake_vision)
    result = await extract_material("a.pdf", "application/pdf", b"%PDF-fake")
    assert result.source == "vision"
    assert result.diagnosis == "骨折"


async def test_pdf_text_model_failure_falls_to_render(monkeypatch) -> None:
    """文本充足但模型提取失败 → 继续走渲染路径，而非直接降级。"""
    from services import materials

    monkeypatch.setattr(
        materials, "_pdf_text", lambda content: "诊断证明 急性阑尾炎 住院费用 15800 元 患者张三" * 2
    )
    monkeypatch.setattr(materials, "_pdf_page_images", lambda c, m: [b"page1"])

    async def fail(text: str) -> None:
        return None

    async def fake_vision(png: bytes) -> dict:
        return {"diagnosis": "阑尾炎"}

    monkeypatch.setattr(materials, "_extract_via_text_model", fail)
    monkeypatch.setattr(materials, "_vision_extract_png", fake_vision)
    result = await extract_material("a.pdf", "application/pdf", b"%PDF-fake")
    assert result.source == "vision"


async def test_pdf_all_paths_fail_falls_back(monkeypatch) -> None:
    """文本与渲染全失败 → Mock 兜底。"""
    from services import materials

    monkeypatch.setattr(materials, "_pdf_text", lambda content: "")
    monkeypatch.setattr(materials, "_pdf_page_images", lambda c, m: [])
    result = await extract_material("a.pdf", "application/pdf", b"%PDF-fake")
    assert result.source == "mock_fallback"
    assert result.file_type == "pdf"
